import pandas as pd
import psycopg
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
import os
from dotenv import load_dotenv

load_dotenv()

def export_attendance():
    connection = psycopg.connect(
        host=os.getenv("DB_HOST"),
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD")
    )

    query = """
    SELECT
        a.attendance_id,
        a.student_id,
        s."NAME" AS name,
        a.attendance_date,
        TO_CHAR(a.attendance_time, 'HH12:MI:SS AM') AS attendance_time,
        a.status
    FROM attendance a
    JOIN students s
    ON a.student_id = s.student_id
    ORDER BY a.attendance_date DESC, a.attendance_time DESC;
    """

    df = pd.read_sql_query(query, connection)

    connection.close()

    df["student_id"] = df["student_id"].astype(str)

    file_name = "attendance.xlsx"

    df.to_excel(
        file_name,
        index=False,
        sheet_name="Attendance"
    )

    workbook = load_workbook(file_name)
    sheet = workbook["Attendance"]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    for row in range(2, sheet.max_row + 1):
        sheet.cell(row, 2).number_format = "@"

        for col in [1, 2, 4, 5, 6]:
            sheet.cell(row, col).alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 20
    sheet.column_dimensions["C"].width = 20
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 20
    sheet.column_dimensions["F"].width = 15

    sheet.row_dimensions[1].height = 25

    sheet.freeze_panes = "A2"

    workbook.save(file_name)

    print("Attendance Excel file created successfully!")

if __name__ == "__main__":
    export_attendance()